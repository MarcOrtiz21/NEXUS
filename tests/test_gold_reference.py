import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from gold_reference import load_reference, parse_implementar


class GoldReferenceTests(unittest.TestCase):
    def _workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "IMPLEMENTAR"
        sheet["A3"] = "Junio Q2"
        sheet["I3"] = "Mayo Q2"
        sheet["Q3"] = "Julio Q3"
        sheet["A17"] = "FMI/3BC"
        sheet["A18"] = "China / PBoC"
        sheet["B18"] = "Compras"
        sheet["G18"] = 1
        sheet["A19"] = "India / NBP"
        sheet["G19"] = 1
        sheet["A20"] = "Polonia / RBI"
        sheet["G20"] = 1
        sheet["A15"] = "RG Real"
        sheet["B15"] = "3,57-3,4 = 0,17"
        sheet["G15"] = -1
        return workbook

    def test_country_reserves_are_not_fmi_or_world_total(self):
        result = parse_implementar(self._workbook()["IMPLEMENTAR"])
        by_cell = {row["source_cell"]: row for row in result["rows"]}
        self.assertEqual(by_cell["A19"]["original_label"], "India / NBP")
        self.assertEqual(by_cell["A19"]["display_label"], "India · RBI")
        self.assertEqual(by_cell["A19"]["scope_id"], "IN")
        self.assertEqual(by_cell["A20"]["display_label"], "Polonia · NBP")
        self.assertEqual(by_cell["A20"]["scope_id"], "PL")
        self.assertEqual(by_cell["A18"]["scope_id"], "CN")
        self.assertEqual(by_cell["A17"]["family"], "ambiguous_header")
        self.assertEqual(by_cell["A17"]["scope_kind"], "unresolved")
        self.assertEqual(by_cell["A19"]["period"], "2026-06")
        self.assertEqual(by_cell["A19"]["legacy_vote"], 1)
        self.assertIsNone(by_cell["A19"]["release_at"])
        self.assertTrue(all(not row["score_eligible"] for row in result["rows"]))

    def test_rg_original_formula_is_preserved_not_converted_to_tips(self):
        result = parse_implementar(self._workbook()["IMPLEMENTAR"])
        rg = next(row for row in result["rows"] if row["source_cell"] == "A15")
        self.assertEqual(rg["family"], "rg_legacy")
        self.assertEqual(rg["date_as_written"], "3,57-3,4 = 0,17")
        self.assertEqual(rg["legacy_vote"], -1)
        self.assertIsNone(rg["release_at"])

    def test_import_reads_only_implementar(self):
        workbook = self._workbook()
        workbook.create_sheet("Hoja 1")["A1"] = "Borrador"
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reference.xlsx"
            workbook.save(path)
            result = load_reference(path)
        self.assertEqual(result["source_sheet"], "IMPLEMENTAR")
        self.assertEqual(result["source_file"], "reference.xlsx")
        self.assertFalse(any(row["original_label"] == "Borrador" for row in result["rows"]))

    def test_missing_implementar_is_rejected(self):
        workbook = Workbook()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "draft.xlsx"
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, "IMPLEMENTAR"):
                load_reference(path)


if __name__ == "__main__":
    unittest.main()
