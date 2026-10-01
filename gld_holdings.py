"""Tenencias físicas de oro del ETF GLD (SPDR Gold Trust).

AVISO DE DATOS — el código de este módulo es parte de NEXUS, pero los datos
que descarga NO lo son. El archivo histórico de GLD pertenece a World Gold
Trust Services, LLC (WGTS). Sus condiciones permiten guardarlo, mostrarlo o
imprimirlo solo para uso personal y no comercial, y prohíben copiarlo,
redistribuirlo, publicarlo o crear obras derivadas sin autorización escrita.

Por eso NEXUS no incluye ni distribuye estos datos: cada usuario los descarga
por su cuenta, bajo su responsabilidad, y quedan solo en ``data/`` (ignorado
por git). No añadas al repositorio filas reales, capturas ni exportaciones de
esta serie; las pruebas usan valores sintéticos.

Condiciones: https://www.spdrgoldshares.com/terms-and-conditions/

Disponibilidad point-in-time: el valor del día T se fecha disponible el
siguiente día hábil de EE. UU. a las 07:00 NYT. El archivo se sobrescribe y no
conserva vintages, así que cada descarga se registra con hash y fecha de
captura, y se cuentan las filas históricas que cambian entre descargas.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import pandas as pd
import requests
from pandas.tseries.holiday import USFederalHolidayCalendar
from pandas.tseries.offsets import CustomBusinessDay

from config import (
    GLD_HOLDINGS_CACHE_FILE,
    GLD_HOLDINGS_CACHE_TTL_SECONDS,
    GLD_HOLDINGS_RAW_DIR,
    GLD_HOLDINGS_RAW_RETENTION,
)


GLD_FUND_ID = "GLD"
GLD_SOURCE = "WGTS:SPDR Gold Trust historical archive"
GLD_ARCHIVE_URL = "https://api.spdrgoldshares.com/api/v1/historical-archive?product=gld&exchange=NYSE&lang=en"
GLD_PAGE_URL = "https://www.spdrgoldshares.com/usa/historical-data/"
GLD_TERMS_URL = "https://www.spdrgoldshares.com/terms-and-conditions/"
GLD_DATA_NOTICE = (
    "Datos de World Gold Trust Services para uso personal y no comercial. "
    "NEXUS no los distribuye: se descargan localmente y no salen de data/."
)
GLD_RELEASE_POLICY = "Día T disponible desde el siguiente día hábil EE. UU. a las 07:00 NYT."

_SHEET_NAME = "US GLD Historical Archive"
_COLUMNS = {
    "date": "date",
    "closing price": "close",
    "ounces of gold per share": "ounces_per_share",
    "nav/share at 10:30am nyt": "nav_per_share",
    "total ounces of gold in the trust": "ounces",
    "tonnes of gold": "tonnes",
    "total net asset value in the trust": "nav_total",
}
_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_CHANGE_WINDOWS = (1, 5, 21, 63)
_PERCENTILE_WINDOW = 756
_STALE_AFTER_DAYS = 5
_REVISION_TOLERANCE_TONNES = 0.005
_EASTERN = ZoneInfo("America/New_York")
_BUSINESS_DAY = CustomBusinessDay(calendar=USFederalHolidayCalendar())


def _number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        pass
    parts = text.split("-")
    if len(parts) != 3 or parts[1][:3].lower() not in _MONTHS:
        return None
    try:
        return date(int(parts[2]), _MONTHS[parts[1][:3].lower()], int(parts[0]))
    except ValueError:
        return None


def release_at_for_session(session: date) -> datetime:
    """Primera hora conservadora en la que el dato de ``session`` es público."""
    release_date = (pd.Timestamp(session) + _BUSINESS_DAY).date()
    return datetime.combine(release_date, time(7, 0), tzinfo=_EASTERN).astimezone(timezone.utc)


def parse_gld_archive(content: bytes) -> list[dict[str, Any]]:
    """Convierte el XLSX del archivo GLD en filas ordenadas por fecha."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook[_SHEET_NAME] if _SHEET_NAME in workbook.sheetnames else None
        if sheet is None:
            raise ValueError("El archivo GLD no contiene la hoja histórica esperada")
        rows = sheet.iter_rows(values_only=True)
        positions: dict[str, int] = {}
        for header in rows:
            labels = [str(cell or "").strip().lower() for cell in header]
            if "date" in labels and "tonnes of gold" in labels:
                positions = {
                    field: labels.index(label)
                    for label, field in _COLUMNS.items() if label in labels
                }
                break
        if "tonnes" not in positions:
            raise ValueError("El archivo GLD no contiene la columna de toneladas")

        parsed: dict[str, dict[str, Any]] = {}
        for raw in rows:
            def cell(field: str) -> Any:
                index = positions.get(field)
                return raw[index] if index is not None and index < len(raw) else None

            session = _parse_date(cell("date"))
            tonnes = _number(cell("tonnes"))
            if session is None or tonnes is None or tonnes <= 0:
                continue
            ounces = _number(cell("ounces"))
            ounces_per_share = _number(cell("ounces_per_share"))
            shares = (
                round(ounces / ounces_per_share)
                if ounces is not None and ounces_per_share else None
            )
            parsed[session.isoformat()] = {
                "date": session.isoformat(),
                "release_at": release_at_for_session(session).isoformat(timespec="seconds"),
                "fund_id": GLD_FUND_ID,
                "tonnes": round(tonnes, 4),
                "ounces": ounces,
                "ounces_per_share": ounces_per_share,
                "shares_outstanding_derived": shares,
                "nav_total_usd": _number(cell("nav_total")),
                "nav_per_share_usd": _number(cell("nav_per_share")),
                "close_usd": _number(cell("close")),
            }
        return [parsed[key] for key in sorted(parsed)]
    finally:
        workbook.close()


def count_revisions(previous: Iterable[dict[str, Any]], current: Iterable[dict[str, Any]]) -> int:
    """Cuenta sesiones ya publicadas cuyo tonelaje cambia entre dos descargas."""
    known = {row.get("date"): _number(row.get("tonnes")) for row in previous}
    revised = 0
    for row in current:
        before = known.get(row.get("date"))
        after = _number(row.get("tonnes"))
        if before is not None and after is not None and abs(before - after) > _REVISION_TOLERANCE_TONNES:
            revised += 1
    return revised


def _read_cache(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    if not isinstance(payload, dict) or not isinstance(payload.get("rows"), list):
        return {}
    return payload


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def _cache_fresh(path: Path) -> bool:
    if not path.exists():
        return False
    age = datetime.now(timezone.utc).timestamp() - path.stat().st_mtime
    return age <= GLD_HOLDINGS_CACHE_TTL_SECONDS


def _store_raw_snapshot(
    content: bytes,
    *,
    digest: str,
    fetched_at: str,
    raw_dir: Path,
    retention: int,
    manifest_entry: dict[str, Any],
) -> None:
    raw_dir.mkdir(parents=True, exist_ok=True)
    stamp = fetched_at[:19].replace(":", "").replace("-", "")
    target = raw_dir / f"gld_archive_{stamp}_{digest[:12]}.xlsx"
    existing = sorted(raw_dir.glob(f"gld_archive_*_{digest[:12]}.xlsx"))
    if not existing:
        target.write_bytes(content)
    with (raw_dir / "manifest.jsonl").open("a", encoding="utf-8") as handle:
        handle.write(json.dumps({**manifest_entry, "file": (existing[0] if existing else target).name}) + "\n")
    snapshots = sorted(raw_dir.glob("gld_archive_*.xlsx"))
    for old in snapshots[: max(0, len(snapshots) - max(1, retention))]:
        old.unlink(missing_ok=True)


def fetch_gld_holdings(
    *,
    refresh: bool = False,
    cache_path: Path = GLD_HOLDINGS_CACHE_FILE,
    raw_dir: Path = GLD_HOLDINGS_RAW_DIR,
    retention: int = GLD_HOLDINGS_RAW_RETENTION,
    session: Any = None,
) -> dict[str, Any]:
    """Descarga el archivo GLD y degrada a la caché local si falla."""
    cached = _read_cache(cache_path)
    if cached and not refresh and _cache_fresh(cache_path):
        return {**cached, "source_status": "CACHE"}

    client = session or requests
    fetched_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        response = client.get(
            GLD_ARCHIVE_URL,
            timeout=30,
            headers={"User-Agent": "NEXUS personal research (local, non-commercial)"},
        )
        response.raise_for_status()
        content = response.content
        rows = parse_gld_archive(content)
        if not rows:
            raise ValueError("El archivo GLD no contiene filas válidas")
    except Exception as exc:
        if cached:
            return {**cached, "source_status": "STALE", "error": str(exc)[:200]}
        return {"rows": [], "source_status": "MISSING", "error": str(exc)[:200]}

    digest = hashlib.sha256(content).hexdigest()
    revised = count_revisions(cached.get("rows") or [], rows)
    snapshot = {
        "fetched_at": fetched_at,
        "sha256": digest,
        "last_modified": str((getattr(response, "headers", None) or {}).get("Last-Modified") or "") or None,
        "rows_count": len(rows),
        "first_date": rows[0]["date"],
        "last_date": rows[-1]["date"],
        "revised_rows_vs_previous": revised,
    }
    try:
        _store_raw_snapshot(
            content, digest=digest, fetched_at=fetched_at, raw_dir=raw_dir,
            retention=retention, manifest_entry=snapshot,
        )
    except OSError as exc:
        snapshot["raw_snapshot_error"] = str(exc)[:200]
    payload = {
        "source": GLD_SOURCE,
        "source_url": GLD_ARCHIVE_URL,
        "terms_url": GLD_TERMS_URL,
        "fund_id": GLD_FUND_ID,
        "snapshot": snapshot,
        "rows": rows,
    }
    _write_json_atomic(cache_path, payload)
    return {**payload, "source_status": "OK"}


def _cutoff_instant(value: date | datetime | str | None) -> datetime:
    if value is None:
        return datetime.now(timezone.utc)
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    return datetime.combine(value, time.max, tzinfo=timezone.utc)


def known_holdings_rows(
    rows: Iterable[dict[str, Any]], cutoff: date | datetime | str | None
) -> list[dict[str, Any]]:
    """Devuelve solo las sesiones publicadas antes del corte."""
    limit = _cutoff_instant(cutoff)
    known = []
    for row in rows:
        try:
            released = datetime.fromisoformat(str(row.get("release_at")))
        except ValueError:
            continue
        if released <= limit:
            known.append(row)
    return sorted(known, key=lambda item: item["date"])


def _change(ordered: list[dict[str, Any]], sessions: int) -> tuple[float | None, float | None]:
    if len(ordered) <= sessions:
        return None, None
    latest = float(ordered[-1]["tonnes"])
    prior = float(ordered[-1 - sessions]["tonnes"])
    pct = (latest / prior - 1) * 100 if prior else None
    return round(latest - prior, 3), (round(pct, 4) if pct is not None else None)


def summarize_gld_holdings(
    source: dict[str, Any] | list[dict[str, Any]],
    *,
    as_of: date | datetime | str | None = None,
    history_sessions: int = 260,
) -> dict[str, Any]:
    """Resume stock y diferencia de stock conocidos en ``as_of``."""
    payload = source if isinstance(source, dict) else {"rows": source, "source_status": "OK"}
    cutoff = _cutoff_instant(as_of)
    base = {
        "fund_id": GLD_FUND_ID,
        "measure": "etf_physical_holdings",
        "unit": "metric_tonnes",
        "change_kind": "difference_in_stock",
        "source": GLD_SOURCE,
        "source_url": GLD_PAGE_URL,
        "terms_url": GLD_TERMS_URL,
        "release_policy": GLD_RELEASE_POLICY,
        "data_notice": GLD_DATA_NOTICE,
        "source_status": payload.get("source_status"),
        "snapshot": payload.get("snapshot"),
        "score_enabled": False,
    }
    ordered = known_holdings_rows(payload.get("rows") or [], cutoff)
    if not ordered:
        return {**base, "status": "MISSING", "history": [], "error": payload.get("error")}

    latest = ordered[-1]
    age_days = (cutoff.date() - date.fromisoformat(latest["date"])).days
    changes = {}
    for window in _CHANGE_WINDOWS:
        tonnes, pct = _change(ordered, window)
        changes[f"change_{window}d_tonnes"] = tonnes
        changes[f"change_{window}d_pct"] = pct

    percentile = None
    if changes["change_21d_pct"] is not None:
        window = ordered[-(_PERCENTILE_WINDOW + 21):]
        history = [
            (float(window[index]["tonnes"]) / float(window[index - 21]["tonnes"]) - 1) * 100
            for index in range(21, len(window)) if float(window[index - 21]["tonnes"])
        ]
        if len(history) >= 252:
            current = changes["change_21d_pct"]
            percentile = round(sum(value <= current for value in history) / len(history) * 100, 1)

    return {
        **base,
        "status": "OK" if age_days <= _STALE_AFTER_DAYS else "STALE",
        "as_of": latest["date"],
        "release_at": latest["release_at"],
        "age_days": age_days,
        "tonnes": latest["tonnes"],
        "ounces": latest.get("ounces"),
        "shares_outstanding_derived": latest.get("shares_outstanding_derived"),
        "shares_note": "Participaciones derivadas: onzas totales / onzas por participación.",
        "nav_total_usd": latest.get("nav_total_usd"),
        **changes,
        "change_21d_percentile_3y": percentile,
        "history": [
            {"date": row["date"], "value": row["tonnes"]}
            for row in (ordered[-history_sessions:] if history_sessions > 0 else [])
        ],
        "error": payload.get("error"),
    }


def holdings_data_fields(summary: dict[str, Any]) -> dict[str, Any]:
    """Campos planos para el modelo, el historial y el backtest."""
    if summary.get("status") not in {"OK", "STALE"}:
        return {}
    fields = {
        "GLD_Holdings_AsOf": summary.get("as_of"),
        "GLD_Holdings_ReleaseAt": summary.get("release_at"),
        "GLD_Holdings_Tonnes": summary.get("tonnes"),
        "GLD_Holdings_1D_Change_Tonnes": summary.get("change_1d_tonnes"),
        "GLD_Holdings_5D_Change_Tonnes": summary.get("change_5d_tonnes"),
        "GLD_Holdings_21D_Change_Tonnes": summary.get("change_21d_tonnes"),
        "GLD_Holdings_5D_Change_Pct": summary.get("change_5d_pct"),
        "GLD_Holdings_21D_Change_Pct": summary.get("change_21d_pct"),
        "GLD_Holdings_63D_Change_Pct": summary.get("change_63d_pct"),
        "GLD_Holdings_21D_Change_Percentile_3Y": summary.get("change_21d_percentile_3y"),
    }
    return {key: value for key, value in fields.items() if value is not None}
