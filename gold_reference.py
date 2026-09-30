"""Importador de referencia, no de observaciones verificadas, para IMPLEMENTAR.

El libro contiene juicios manuales y fechas sin hora de publicación comprobada.
Nunca se envían estas filas al histórico point-in-time ni al score.
"""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any


MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5,
    "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9,
}
BLOCK_ROWS = ((3, 20), (24, 41), (44, 61))
BLOCK_COLUMNS = (1, 9, 17)
RESERVE_LABELS = {
    "China / PBoC": ("China · PBoC", "CN"),
    "India / NBP": ("India · RBI", "IN"),
    "Polonia / RBI": ("Polonia · NBP", "PL"),
}


def _raw(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return None if value in {"", "/"} else value
    return value


def _period(header: Any) -> str | None:
    match = re.fullmatch(r"([A-Za-zÁÉÍÓÚáéíóú]+)\s+Q[1-4]", str(header or "").strip())
    if not match:
        return None
    month = MONTHS.get(match.group(1).lower())
    return f"2026-{month:02d}" if month else None


def parse_implementar(sheet: Any) -> dict[str, Any]:
    """Preserva celdas y votos originales, corrigiendo solo el rótulo de país.

    ``sheet`` puede ser una hoja openpyxl o un doble de pruebas con ``cell``.
    Una fecha de la columna B/J/R es solo texto de referencia: no release_at.
    """
    rows: list[dict[str, Any]] = []
    warnings: list[str] = []
    from openpyxl.utils import get_column_letter

    for first_row, last_row in BLOCK_ROWS:
        for first_col in BLOCK_COLUMNS:
            period = _period(sheet.cell(first_row, first_col).value)
            if period is None:
                warnings.append(f"Cabecera no reconocida: {get_column_letter(first_col)}{first_row}")
                continue
            for row_index in range(first_row + 1, last_row + 1):
                cells = [sheet.cell(row_index, first_col + offset) for offset in range(7)]
                original = _raw(cells[0].value)
                if original is None:
                    continue
                corrected = RESERVE_LABELS.get(original)
                if corrected:
                    family, scope_kind, scope_id = "official_reserve_stock", "country", corrected[1]
                elif original == "FMI/3BC":
                    family, scope_kind, scope_id = "ambiguous_header", "unresolved", None
                elif str(original).startswith("RG"):
                    family, scope_kind, scope_id = "rg_legacy", "unresolved", None
                elif original in {"WGC", "GMC(Mensual)"}:
                    family, scope_kind, scope_id = "editorial_reference", "unresolved", None
                else:
                    family, scope_kind, scope_id = "macro_reference", "unresolved", None
                rows.append({
                    "period": period,
                    "source_sheet": "IMPLEMENTAR",
                    "source_cell": f"{get_column_letter(first_col)}{row_index}",
                    "original_label": original,
                    "display_label": corrected[0] if corrected else original,
                    "family": family,
                    "scope_kind": scope_kind,
                    "scope_id": scope_id,
                    "date_as_written": _raw(cells[1].value),
                    "expected_as_written": _raw(cells[2].value),
                    "actual_as_written": _raw(cells[3].value),
                    "versus_expected_as_written": _raw(cells[4].value),
                    "conclusion_as_written": _raw(cells[5].value),
                    "legacy_vote": _raw(cells[6].value),
                    "release_at": None,
                    "quality": "UNVERIFIED_REFERENCE",
                    "score_eligible": False,
                })
    return {
        "kind": "gold_workbook_reference",
        "source_sheet": "IMPLEMENTAR",
        "score_eligible": False,
        "rows": rows,
        "warnings": warnings,
        "note": "Votos y fechas transcritos sin validar; reservas por país, no agregado mundial ni FMI/3BC.",
    }


def load_reference(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "IMPLEMENTAR" not in workbook.sheetnames:
            raise ValueError("Falta la hoja IMPLEMENTAR")
        result = parse_implementar(workbook["IMPLEMENTAR"])
        result["source_file"] = path.name
        return result
    finally:
        workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Inspeccionar IMPLEMENTAR como referencia sin score")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    print(json.dumps(load_reference(args.workbook), ensure_ascii=False, indent=2))
